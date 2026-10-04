import io

import openpyxl
from fastapi.testclient import TestClient
from sqlmodel import SQLModel, Session, create_engine
from sqlmodel.pool import StaticPool

from .database import get_session
from .importers import parse_rows
from .main import app

engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
SQLModel.metadata.create_all(engine)
app.dependency_overrides[get_session] = lambda: Session(engine)
client = TestClient(app)

# Auth: expenses are scoped to a logged-in user; TestClient persists the
# session cookie set on signup across the rest of these requests.
assert client.get("/api/expenses").status_code == 401
assert client.post("/api/auth/signup", json={"username": "asha", "password": "s3cret"}).json() == {"username": "asha"}
assert client.post("/api/auth/signup", json={"username": "asha", "password": "other"}).status_code == 400
assert client.post("/api/auth/login", json={"username": "asha", "password": "wrong"}).status_code == 401
assert client.get("/api/auth/me").json() == {"username": "asha"}

other = TestClient(app)
other.post("/api/auth/signup", json={"username": "ravi", "password": "s3cret"})
other.post("/api/expenses", json={"date": "2026-10-01", "description": "ravi's lunch", "amount": 99, "category": "Food & Dining"})
assert [e["description"] for e in client.get("/api/expenses").json()] == []  # asha can't see ravi's expenses

created = client.post("/api/expenses", json={
    "date": "2026-10-01", "description": "lunch", "amount": 150, "category": "Food & Dining",
}).json()
assert created["id"] is not None

listed = client.get("/api/expenses").json()
assert len(listed) == 1 and listed[0]["description"] == "lunch"

assert client.delete(f"/api/expenses/{created['id']}").status_code == 204
assert client.get("/api/expenses").json() == []
assert client.delete(f"/api/expenses/{created['id']}").status_code == 404

# parse_rows: single signed "Amount" column — negative is a spend, positive is a credit to skip
assert parse_rows([
    {"Date": "04/10/2026", "Description": "Swiggy dinner", "Amount": "-450.00"},
    {"Date": "05/10/2026", "Description": "Salary", "Amount": "20000"},
]) == [{"date": "2026-10-04", "description": "Swiggy dinner", "amount": 450.0}]

# parse_rows: separate Debit/Credit columns — blank debit cell means it was a credit row
assert parse_rows([
    {"Date": "2026-10-05", "Narration": "Bus fare", "Debit": "40", "Credit": ""},
    {"Date": "2026-10-06", "Narration": "Refund", "Debit": "", "Credit": "100"},
]) == [{"date": "2026-10-05", "description": "Bus fare", "amount": 40.0}]

# CSV import via the API: dedups on re-upload, categorizes from description
csv_bytes = b"Date,Description,Amount\n04/10/2026,Swiggy dinner,-450.00\n05/10/2026,Salary,20000\n"
first = client.post("/api/expenses/import", files={"file": ("statement.csv", csv_bytes, "text/csv")}).json()
assert first == {"total_rows": 1, "imported": 1, "skipped_duplicates": 0}
again = client.post("/api/expenses/import", files={"file": ("statement.csv", csv_bytes, "text/csv")}).json()
assert again == {"total_rows": 1, "imported": 0, "skipped_duplicates": 1}
imported_expense = client.get("/api/expenses").json()[0]
assert imported_expense["category"] == "Food & Dining" and imported_expense["source"] == "csv"
client.delete(f"/api/expenses/{imported_expense['id']}")

# xlsx import (Paytm-style UPI statement): "Transaction Details" as the
# description column, a non-transactions sheet to skip past.
wb = openpyxl.Workbook()
wb.active.title = "Summary"
wb.active.append(["Money Paid", "-100.00"])
ws = wb.create_sheet("Passbook Payment History")
ws.append(["Date", "Time", "Transaction Details", "Amount", "Tags"])
# description alone has no category keywords; the app's own tag should win
ws.append(["03/10/2026", "20:04:04", "Paid to Sohani Wo Dhanna Ra", "-40.00", "#🥘 Food"])
ws.append(["02/10/2026", "12:54:18", "Received from Geethika", "+140.00", "#💵 Money Received"])
xlsx_buf = io.BytesIO()
wb.save(xlsx_buf)
xlsx_result = client.post(
    "/api/expenses/import",
    files={"file": ("statement.xlsx", xlsx_buf.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
).json()
assert xlsx_result == {"total_rows": 1, "imported": 1, "skipped_duplicates": 0}
xlsx_expense = client.get("/api/expenses").json()[0]
assert xlsx_expense["source"] == "xlsx" and xlsx_expense["amount"] == 40.0
assert xlsx_expense["category"] == "Food & Dining"
client.delete(f"/api/expenses/{xlsx_expense['id']}")

# xlsx import: two genuinely separate transactions with the same date, payee,
# and amount must NOT be deduped against each other when a UPI Ref No. is present
wb2 = openpyxl.Workbook()
ws2 = wb2.create_sheet("Passbook Payment History")
del wb2["Sheet"]
ws2.append(["Date", "Transaction Details", "Amount", "UPI Ref No."])
ws2.append(["03/10/2026", "Paid to Pampati Ismail", "-13.00", "111111111111"])
ws2.append(["03/10/2026", "Paid to Pampati Ismail", "-13.00", "222222222222"])
xlsx2_buf = io.BytesIO()
wb2.save(xlsx2_buf)
xlsx2_result = client.post(
    "/api/expenses/import",
    files={"file": ("statement2.xlsx", xlsx2_buf.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
).json()
assert xlsx2_result == {"total_rows": 2, "imported": 2, "skipped_duplicates": 0}
# re-importing the same file is still deduped via the ref no.
xlsx2_again = client.post(
    "/api/expenses/import",
    files={"file": ("statement2.xlsx", xlsx2_buf.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
).json()
assert xlsx2_again == {"total_rows": 2, "imported": 0, "skipped_duplicates": 2}
for e in client.get("/api/expenses").json():
    client.delete(f"/api/expenses/{e['id']}")

assert client.post("/api/auth/logout").json() == {"ok": True}
assert client.get("/api/expenses").status_code == 401

print("all backend self-checks passed")
